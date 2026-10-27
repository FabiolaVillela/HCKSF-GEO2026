"""Results report: per-bar force envelope and member checks, per-node deflection and support type.

All stages are solved with `assembly.solve_stage`; each bar and node keeps the worst value it
reaches during the assembly and the stage where it happens. Stages without a solution
(mechanism or error) are listed separately and never contribute values.

Member checks come from `aisc.check_result` (preliminary ANSI/AISC 360-22 LRFD, round steel HSS,
1.4D) when a material grade is selected. Without a grade the bars are reported as "not checked";
for aluminum AISC does not apply.

`build_report` returns language-neutral data (codes, numbers). `report_xlsx(report, language)`
writes the workbook in English ('en') or Spanish ('es'): sheet names, headers, values and
engine messages are all translated there.
"""
from __future__ import annotations

from io import BytesIO
import math

import numpy as np
import pandas as pd

from .aisc import check_result
from .assembly import HSS, solve_stage, stage_plan, validate_geometry
from .design_materials import design_material

LANGUAGES = ('en', 'es')


# Engine messages come in Spanish; English versions (same as the simulator's translations).
ENGINE_MESSAGES_EN = {
    'No hay barras instaladas en esta etapa.': 'No bars are installed in this stage.',
    'Equilibrio elástico lineal calculado.': 'Linear elastic equilibrium calculated.',
    'Modelo articulado inestable: traslaciones sin rigidez.': 'Unstable pinned model: translations without stiffness.',
    'Modelo articulado inestable: mecanismo o matriz mal condicionada.': 'Unstable pinned model: mechanism or ill-conditioned matrix.',
    'Modelo inestable: GDL sin rigidez': 'Unstable model: degrees of freedom without stiffness',
    'Modelo inestable o excesivamente mal condicionado': 'Unstable or severely ill-conditioned model',
    'Residual de solución excesivo': 'Excessive solution residual',
    'Residual de solución excesivo.': 'Excessive solution residual.',
}

TEXT = {
    'sheets': {'en': ['Summary', 'Bars', 'Nodes', 'Stages'], 'es': ['Resumen', 'Barras', 'Nodos', 'Etapas']},
    'summary_header': {'en': ['Item', 'Value'], 'es': ['Concepto', 'Valor']},
    'summary_items': {
        'material': {'en': 'Material', 'es': 'Material'},
        'profile': {'en': 'Profile', 'es': 'Perfil'},
        'connections': {'en': 'Connections', 'es': 'Conexiones'},
        'default_support': {'en': 'Default support', 'es': 'Apoyo por defecto'},
        'bars_per_stage': {'en': 'Bars per stage', 'es': 'Barras por etapa'},
        'installer_load_kg': {'en': 'Installer load (kg)', 'es': 'Carga del instalador (kg)'},
        'total_stages': {'en': 'Total stages', 'es': 'Etapas totales'},
        'solved_stages': {'en': 'Solved stages', 'es': 'Etapas resueltas'},
        'stages_without_solution': {'en': 'Stages without solution', 'es': 'Etapas sin solución'},
        'total_bars': {'en': 'Total bars', 'es': 'Barras totales'},
        'total_nodes': {'en': 'Total nodes', 'es': 'Nodos totales'},
        'member_checks': {'en': 'Member checks', 'es': 'Revisiones de elementos'},
        'material_grade': {'en': 'Material grade', 'es': 'Grado de material'},
        'fy_fu': {'en': 'Fy / Fu (MPa)', 'es': 'Fy / Fu (MPa)'},
        'grade_source': {'en': 'Strength source', 'es': 'Fuente de resistencias'},
        'design_basis': {'en': 'Design basis', 'es': 'Bases de diseño'},
        'check_scope': {'en': 'Scope of the checks', 'es': 'Alcance de las revisiones'},
        'max_utilization': {'en': 'Max utilization (all bars, all stages)', 'es': 'Utilización máx. (todas las barras y etapas)'},
        'bars_exceeding': {'en': 'Bars with utilization > 1', 'es': 'Barras con utilización > 1'},
        'load_case': {'en': 'Load case', 'es': 'Caso de carga'},
        'note': {'en': 'Note', 'es': 'Nota'},
    },
    'bars': {
        'bar_id': {'en': 'Bar', 'es': 'Barra'},
        'node_i': {'en': 'Node i', 'es': 'Nodo i'},
        'node_j': {'en': 'Node j', 'es': 'Nodo j'},
        'length_m': {'en': 'Length (m)', 'es': 'Longitud (m)'},
        'installed_stage': {'en': 'Installed at stage', 'es': 'Etapa de instalación'},
        'max_tension_kN': {'en': 'Max tension (kN)', 'es': 'Tracción máx. (kN)'},
        'tension_stage': {'en': 'Stage of max tension', 'es': 'Etapa de tracción máx.'},
        'max_compression_kN': {'en': 'Max compression (kN)', 'es': 'Compresión máx. (kN)'},
        'compression_stage': {'en': 'Stage of max compression', 'es': 'Etapa de compresión máx.'},
        'check_status': {'en': 'Member check', 'es': 'Revisión del elemento'},
        'max_utilization': {'en': 'Max utilization', 'es': 'Utilización máx.'},
        'governing_check': {'en': 'Governing check (AISC 360-22)', 'es': 'Revisión que rige (AISC 360-22)'},
        'limit_state': {'en': 'Limit state', 'es': 'Estado límite'},
        'check_x_m': {'en': 'Position along bar (m)', 'es': 'Posición en la barra (m)'},
        'utilization_stage': {'en': 'Stage of max utilization', 'es': 'Etapa de utilización máx.'},
        'check_note': {'en': 'Check note', 'es': 'Nota de la revisión'},
    },
    'nodes': {
        'node_id': {'en': 'Node', 'es': 'Nodo'},
        'x_m': {'en': 'X (m)', 'es': 'X (m)'},
        'y_m': {'en': 'Y (m)', 'es': 'Y (m)'},
        'z_m': {'en': 'Z (m)', 'es': 'Z (m)'},
        'support_type': {'en': 'Support type', 'es': 'Tipo de apoyo'},
        'max_deflection_mm': {'en': 'Max deflection (mm)', 'es': 'Desplazamiento máx. (mm)'},
        'deflection_stage': {'en': 'Stage of max deflection', 'es': 'Etapa de desplazamiento máx.'},
        'ux_mm': {'en': 'ux (mm)', 'es': 'ux (mm)'},
        'uy_mm': {'en': 'uy (mm)', 'es': 'uy (mm)'},
        'uz_mm': {'en': 'uz (mm)', 'es': 'uz (mm)'},
    },
    'stages': {
        'stage': {'en': 'Stage', 'es': 'Etapa'},
        'progress_pct': {'en': 'Progress (%)', 'es': 'Avance (%)'},
        'active_bars': {'en': 'Installed bars', 'es': 'Barras instaladas'},
        'status': {'en': 'Status', 'es': 'Estado'},
        'message': {'en': 'Engine message', 'es': 'Mensaje del motor'},
        'max_tension_kN': {'en': 'Max tension (kN)', 'es': 'Tracción máx. (kN)'},
        'max_compression_kN': {'en': 'Max compression (kN)', 'es': 'Compresión máx. (kN)'},
        'max_deflection_mm': {'en': 'Max deflection (mm)', 'es': 'Desplazamiento máx. (mm)'},
        'max_utilization': {'en': 'Max utilization', 'es': 'Utilización máx.'},
        'governing_bar': {'en': 'Governing bar', 'es': 'Barra que rige'},
    },
    'support': {'empotrado': {'en': 'Fixed', 'es': 'Empotrado'}, 'articulado': {'en': 'Pinned', 'es': 'Articulado'},
                'free': {'en': 'Free', 'es': 'Libre'}},
    'connection': {'empotrada': {'en': 'Rigid', 'es': 'Empotradas'}, 'articulada': {'en': 'Pinned', 'es': 'Articuladas'}},
    'material': {'Acero': {'en': 'Steel', 'es': 'Acero'}, 'Aluminio': {'en': 'Aluminum', 'es': 'Aluminio'}},
    'status': {'stable': {'en': 'Solved', 'es': 'Resuelta'}, 'unstable': {'en': 'Mechanism (no solution)', 'es': 'Mecanismo (sin solución)'},
               'empty': {'en': 'No bars', 'es': 'Sin barras'}, 'error': {'en': 'Error', 'es': 'Error'}},
    'check_status': {'within': {'en': 'Within included checks (≤ 1)', 'es': 'Dentro de las revisiones incluidas (≤ 1)'},
                     'exceeds': {'en': 'Exceeds (> 1)', 'es': 'Excede (> 1)'},
                     'not_checked': {'en': 'Not checked', 'es': 'Sin revisar'},
                     'not_applicable': {'en': 'Not applicable', 'es': 'No aplica'},
                     'unavailable': {'en': 'Unavailable', 'es': 'No disponible'}},
    'limit_state': {'tension': {'en': 'Tension', 'es': 'Tracción'},
                    'compression': {'en': 'Compression', 'es': 'Compresión'},
                    'flexotension': {'en': 'Flexure + tension', 'es': 'Flexotracción'},
                    'flexocompression': {'en': 'Flexure + compression', 'es': 'Flexocompresión'},
                    'biaxial flexure (conservative sum)': {'en': 'Biaxial flexure (conservative sum)',
                                                           'es': 'Flexión biaxial (suma conservadora)'}},
    'checks_mode': {'aisc': {'en': 'Preliminary ANSI/AISC 360-22 LRFD (round steel HSS)',
                             'es': 'Preliminares ANSI/AISC 360-22 LRFD (HSS redondo de acero)'},
                    'no_grade': {'en': 'Not checked: no material grade selected',
                                 'es': 'Sin revisar: no se seleccionó un grado de material'},
                    'aluminum': {'en': 'Not applicable: AISC 360 does not apply to aluminum',
                                 'es': 'No aplica: AISC 360 no aplica al aluminio'},
                    'installer': {'en': 'Not checked: the 1.4D check accepts dead load only (set installer load to 0)',
                                  'es': 'Sin revisar: la revisión 1.4D solo admite carga muerta (pon la carga del instalador en 0)'}},
    'bar_note': {'no_grade': {'en': 'Select a material grade to check this bar.', 'es': 'Selecciona un grado de material para revisar esta barra.'},
                 'aluminum': {'en': 'AISC 360 does not apply to aluminum.', 'es': 'AISC 360 no aplica al aluminio.'},
                 'installer': {'en': 'Installer load must be 0 for the 1.4D check.', 'es': 'La carga del instalador debe ser 0 para la revisión 1.4D.'},
                 'never_solved': {'en': 'Not part of any solved stage.', 'es': 'No forma parte de ninguna etapa resuelta.'},
                 'kl_r': {'en': 'KL/r > 200 (E2 preferred slenderness limit exceeded).', 'es': 'KL/r > 200 (se excede el límite de esbeltez recomendado en E2).'},
                 'out_of_scope': {'en': 'Section outside the scope of the check.', 'es': 'Sección fuera del alcance de la revisión.'}},
    'design_basis': {'en': '1.4D · t = 0.93 t nominal · K = 1 · L = node-to-node · An = Ag · U = 1 · first-order, independent stage',
                     'es': '1.4D · t = 0.93 t nominal · K = 1 · L = entre nodos · An = Ag · U = 1 · primer orden, etapa independiente'},
    'check_scope': {'en': 'Preliminary member checks only (axial, flexure and H1 interaction). Shear, torsion, connections and '
                          'Chapter C stability are not checked. Not a design approval. See docs/AISC.md.',
                    'es': 'Solo revisiones preliminares de elementos (axial, flexión e interacción H1). No se revisan cortante, '
                          'torsión, conexiones ni la estabilidad del capítulo C. No es una aprobación de diseño. Ver docs/AISC.md.'},
    'none': {'en': 'none', 'es': 'ninguna'},
    'self_weight': {'en': 'Self-weight of installed bars', 'es': 'Peso propio de las barras instaladas'},
    'installer': {'en': ' + installer {kg} kg', 'es': ' + instalador {kg} kg'},
    'note': {'en': 'Envelope over all solved stages (independent linear-elastic equilibrium per stage). '
                   'Pinned connections can give near-mechanism stages with unrealistic values (issue #14).',
             'es': 'Envolvente de todas las etapas resueltas (equilibrio elástico lineal independiente por etapa). '
                   'Con conexiones articuladas puede haber etapas casi-mecanismo con valores no realistas (incidencia #14 del repositorio).'},
}


def _support_code(node_id, support_nodes, config):
    if node_id not in support_nodes:
        return 'free'
    overrides = config.get('support_overrides', {}) or {}
    return overrides.get(str(node_id), overrides.get(node_id, config.get('support', 'empotrado')))


def _checks_mode(material, config, design):
    """'aisc' when the checks can run; otherwise the reason they don't."""
    if material != 'Acero':
        return 'aluminum'
    if not design or not design.get('grade'):
        return 'no_grade'
    if config.get('worker_kg', 0):
        return 'installer'
    return 'aisc'


def build_report(geometry, config=None, design=None):
    """Solve every stage and return language-neutral {'summary', 'stages', 'bars', 'nodes'}.

    `design` is the material-grade selection used by `aisc.check_result`
    ({'grade': 'a500-21-c'} or {'grade': 'custom', 'Fy_MPa', 'Fu_MPa', 'source'}).
    """
    config, design = dict(config or {}), dict(design or {})
    nodes, bars, support_nodes = validate_geometry(geometry)
    plan = stage_plan(geometry, config.get('bars_per_step', 2))
    material = config.get('material', 'Acero')
    profile = config.get('profile', list(HSS)[2])
    mode = _checks_mode(material, config, design)
    grade = design_material(material, design) if mode == 'aisc' else None  # raises on an invalid selection
    lengths = {b['id']: math.dist(nodes[b['i']], nodes[b['j']]) for b in bars}
    bar_rows = {b['id']: dict(bar_id=b['id'], node_i=b['i'], node_j=b['j'], length_m=lengths[b['id']],
                              installed_stage=None, max_tension_kN=0.0, tension_stage=None,
                              max_compression_kN=0.0, compression_stage=None,
                              check_status='not_checked' if mode != 'aluminum' else 'not_applicable',
                              max_utilization=None, governing_check=None, limit_state=None, check_x_m=None,
                              utilization_stage=None, check_note=mode if mode != 'aisc' else 'never_solved')
                for b in bars}
    node_rows = {nid: dict(node_id=nid, x_m=p[0], y_m=p[1], z_m=p[2],
                           support_type=_support_code(nid, support_nodes, config),
                           max_deflection_mm=None, deflection_stage=None, ux_mm=None, uy_mm=None, uz_mm=None)
                 for nid, p in nodes.items()}
    stages = []
    for info in plan:
        for bid in info['new_bar_ids']:
            bar_rows[bid]['installed_stage'] = info['stage']
        r = solve_stage(geometry, config, info['stage'])
        stage_row = dict(stage=info['stage'], progress_pct=info['progress_pct'], active_bars=info['active_count'],
                         status=r['status'], message=r['message'],
                         max_tension_kN=r['max_tension_kn'], max_compression_kN=r['max_compression_kn'],
                         max_deflection_mm=r['max_displacement_mm'], max_utilization=None, governing_bar=None)
        stages.append(stage_row)
        if r['status'] != 'stable':
            continue
        for e in r['elements']:
            row = bar_rows[e['id']]
            axial = [s['N_Vy_Vz_T_My_Mz'][0] for s in e['stations']]
            tension, compression = max(0.0, max(axial)) / 1000, max(0.0, -min(axial)) / 1000
            if tension > row['max_tension_kN']:
                row.update(max_tension_kN=tension, tension_stage=info['stage'])
            if compression > row['max_compression_kN']:
                row.update(max_compression_kN=compression, compression_stage=info['stage'])
        if mode == 'aisc':
            checked = check_result(r, design)
            stage_row.update(max_utilization=checked['max_utilization'],
                             governing_bar=(checked.get('governing') or {}).get('member_id'))
            for m in checked['members']:
                row = bar_rows[m['id']]
                if m['utilization'] is None:
                    if row['max_utilization'] is None:
                        row.update(check_status='unavailable', check_note='out_of_scope')
                    continue
                if row['max_utilization'] is None or m['utilization'] > row['max_utilization']:
                    row.update(max_utilization=m['utilization'], governing_check=m['equation'],
                               limit_state=m['limit_state'], check_x_m=m['x_m'], utilization_stage=info['stage'],
                               check_status='within' if m['utilization'] <= 1 else 'exceeds',
                               check_note='kl_r' if m.get('warning') else None)
        for d in r['displacements']:
            u = np.array(d['values'][:3]) * 1000
            row = node_rows[d['node_id']]
            norm = float(np.linalg.norm(u))
            if row['max_deflection_mm'] is None or norm > row['max_deflection_mm']:
                row.update(max_deflection_mm=norm, deflection_stage=info['stage'],
                           ux_mm=float(u[0]), uy_mm=float(u[1]), uz_mm=float(u[2]))
    unsolved = [s['stage'] for s in stages if s['status'] not in ('stable', 'empty')]
    utilizations = [b['max_utilization'] for b in bar_rows.values() if b['max_utilization'] is not None]
    summary = dict(
        material=material, profile=profile, connections=config.get('connection', 'empotrada'),
        default_support=config.get('support', 'empotrado'), bars_per_stage=config.get('bars_per_step', 2),
        installer_load_kg=config.get('worker_kg', 0), total_stages=len(stages),
        solved_stages=sum(s['status'] == 'stable' for s in stages), stages_without_solution=unsolved,
        total_bars=len(bars), total_nodes=len(nodes), checks_mode=mode, grade=grade,
        max_utilization=max(utilizations) if utilizations else None,
        bars_exceeding=sum(u > 1 for u in utilizations))
    return dict(summary=summary, stages=stages, bars=list(bar_rows.values()), nodes=list(node_rows.values()))


def _t(group, code, language):
    entry = TEXT[group].get(code)
    return entry[language] if entry else code


def _localized_tables(report, language):
    s = report['summary']
    load = TEXT['self_weight'][language] + (TEXT['installer'][language].format(kg=s['installer_load_kg']) if s['installer_load_kg'] else '')
    rows = dict(
        material=_t('material', s['material'], language), profile=s['profile'],
        connections=_t('connection', s['connections'], language),
        default_support=_t('support', s['default_support'], language),
        bars_per_stage=s['bars_per_stage'], installer_load_kg=s['installer_load_kg'],
        total_stages=s['total_stages'], solved_stages=s['solved_stages'],
        stages_without_solution=', '.join(map(str, s['stages_without_solution'])) or TEXT['none'][language],
        total_bars=s['total_bars'], total_nodes=s['total_nodes'],
        member_checks=_t('checks_mode', s['checks_mode'], language))
    if s['grade']:
        g = s['grade']
        label = g['label'] if language == 'en' else g['label'].replace(' Grade ', ' Grado ')
        rows.update(material_grade=label if g.get('id') != 'custom' else {'en': 'Custom', 'es': 'Personalizado'}[language],
                    fy_fu=f"{g['Fy_MPa']:g} / {g['Fu_MPa']:g}", grade_source=g['source'],
                    design_basis='ANSI/AISC 360-22 LRFD · ' + TEXT['design_basis'][language],
                    check_scope=TEXT['check_scope'][language],
                    max_utilization=s['max_utilization'], bars_exceeding=s['bars_exceeding'])
    rows.update(load_case=load, note=TEXT['note'][language])
    summary = pd.DataFrame([(TEXT['summary_items'][k][language], v) for k, v in rows.items()],
                           columns=TEXT['summary_header'][language])
    bars = pd.DataFrame(report['bars'], columns=list(TEXT['bars']))
    bars['check_status'] = bars['check_status'].map(lambda c: _t('check_status', c, language))
    bars['limit_state'] = bars['limit_state'].map(lambda c: _t('limit_state', c, language) if isinstance(c, str) else c)
    bars['check_note'] = bars['check_note'].map(lambda c: _t('bar_note', c, language) if isinstance(c, str) else c)
    nodes = pd.DataFrame(report['nodes'], columns=list(TEXT['nodes']))
    nodes['support_type'] = nodes['support_type'].map(lambda c: _t('support', c, language))
    stages = pd.DataFrame(report['stages'], columns=list(TEXT['stages']))
    stages['status'] = stages['status'].map(lambda c: _t('status', c, language))
    if language == 'en':
        stages['message'] = stages['message'].map(lambda m: ENGINE_MESSAGES_EN.get(m, m))
    if not s['grade']:  # no checks ran: drop the empty per-stage utilization columns
        stages = stages.drop(columns=['max_utilization', 'governing_bar'])
    for frame, group in ((bars, 'bars'), (nodes, 'nodes'), (stages, 'stages')):
        frame.columns = [TEXT[group][c][language] for c in frame.columns]
    return [summary, bars, nodes, stages]


def report_xlsx(report, language='en'):
    """Excel workbook (bytes) in English ('en') or Spanish ('es')."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    if language not in LANGUAGES:
        raise ValueError(f'Language must be one of {LANGUAGES}.')
    names = TEXT['sheets'][language]
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        for name, frame in zip(names, _localized_tables(report, language)):
            frame.to_excel(writer, sheet_name=name, index=False)
        book = writer.book
        for sheet in book.worksheets:
            sheet.freeze_panes = 'A2'
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for idx, column in enumerate(sheet.columns, start=1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
                sheet.column_dimensions[get_column_letter(idx)].width = min(max(10, width + 2), 60)
        bars = book[names[1]]
        col = [c.value for c in bars[1]].index(TEXT['bars']['max_utilization'][language]) + 1
        red = PatternFill('solid', fgColor='FECACA')
        for row in bars.iter_rows(min_row=2, min_col=col, max_col=col):
            if isinstance(row[0].value, (int, float)) and row[0].value > 1:
                row[0].fill = red
        summary = book[names[0]]
        checks_label = TEXT['summary_items']['member_checks'][language]
        for row in summary.iter_rows(min_row=2):
            if row[0].value == checks_label and report['summary']['checks_mode'] != 'aisc':
                row[1].font = Font(bold=True, color='B91C1C')
    return buffer.getvalue()
